"""Exportable evidence appendix — citations, SME critique, benchmarks, ingestion."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any, Dict, List

from docx import Document

from app.config import OUTPUT_DIR
from app.services.narrative_composer import build_citations, provenance_banner_text
from app.utils.inr_format import format_money


def _fmt(amount: float, currency: str) -> str:
    return format_money(float(amount or 0.0), currency)


def build_evidence_appendix(
    analysis: Dict[str, Any],
    *,
    company_name: str = "Client",
    narrative_provenance: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    outputs = analysis.get("skill_outputs", {}) or {}
    currency = str(analysis.get("reporting_currency") or "INR")
    provenance = narrative_provenance or {"mode": "deterministic", "llm_skip_reason": "export_no_llm"}
    citations = build_citations(outputs, currency=currency)

    sme = outputs.get("sme-critique", {}) or {}
    critiques = sme.get("initiative_critiques") or []
    sme_rows: List[Dict[str, Any]] = []
    for row in critiques:
        if not isinstance(row, dict):
            continue
        sme_rows.append({
            "category": row.get("category_name") or row.get("category_id"),
            "lever": row.get("lever_name") or row.get("lever"),
            "saving_3yr": _fmt(float(row.get("modelled_saving_3yr") or 0), currency),
            "maturity": row.get("evidence_maturity"),
            "verdict": row.get("sme_verdict"),
            "double_count": bool(row.get("double_count_risk")),
            "gaps": "; ".join(str(g) for g in (row.get("gaps") or [])[:2]),
        })

    peer = outputs.get("peer-benchmarker", {}) or {}
    dataset = peer.get("benchmark_dataset") or peer.get("benchmark_metadata") or {}
    vintage = str(dataset.get("vintage_date") or dataset.get("vintage") or "")
    benchmark_rows: List[Dict[str, Any]] = []
    for row in peer.get("comparisons") or peer.get("benchmark_gaps") or []:
        if not isinstance(row, dict):
            continue
        benchmark_rows.append({
            "category": row.get("category_name") or row.get("category_id"),
            "actual_pct": row.get("actual_pct_of_revenue"),
            "p50_pct": row.get("benchmark_p50_pct") or row.get("benchmark_target_pct"),
            "band": row.get("percentile_band"),
            "gap": _fmt(float(row.get("estimated_saving_amount") or 0), currency),
            "sample_n": row.get("benchmark_sample_n") or dataset.get("sample_n"),
            "vintage": vintage,
        })

    ingestion = analysis.get("ingestion_summary") or analysis.get("ingestion_report") or {}

    return {
        "type": "evidence_appendix",
        "generated_on": str(date.today()),
        "company_name": company_name,
        "currency": currency,
        "narrative_provenance": provenance,
        "narrative_provenance_banner": provenance_banner_text(provenance),
        "citations": [c.model_dump() for c in citations],
        "sme_critique": sme_rows,
        "benchmark_rows": benchmark_rows,
        "ingestion_summary": ingestion,
        "critique_summary": sme.get("critique_summary") or {},
    }


def export_evidence_appendix_docx(appendix: Dict[str, Any], filename: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    doc = Document()
    doc.add_heading(f"Evidence Appendix — {appendix.get('company_name', 'Client')}", level=0)
    banner = appendix.get("narrative_provenance_banner")
    if banner:
        p = doc.add_paragraph(str(banner))
        p.italic = True
    doc.add_paragraph(f"Generated {appendix.get('generated_on')} · Currency {appendix.get('currency')}")

    doc.add_heading("Claim citations", level=1)
    cites = appendix.get("citations") or []
    if cites:
        tbl = doc.add_table(rows=1 + len(cites), cols=4)
        tbl.style = "Table Grid"
        for j, h in enumerate(["Skill", "Metric", "Value", "Rows"]):
            tbl.rows[0].cells[j].text = h
        for i, row in enumerate(cites):
            tbl.rows[i + 1].cells[0].text = str(row.get("source_skill") or "")
            tbl.rows[i + 1].cells[1].text = str(row.get("metric") or "")
            tbl.rows[i + 1].cells[2].text = str(row.get("value") or "")
            tbl.rows[i + 1].cells[3].text = str(row.get("row_count") or "")
    else:
        doc.add_paragraph("No citations recorded.")

    doc.add_heading("SME qualification", level=1)
    sme_rows = appendix.get("sme_critique") or []
    if sme_rows:
        tbl = doc.add_table(rows=1 + len(sme_rows), cols=6)
        tbl.style = "Table Grid"
        headers = ["Category", "Lever", "3yr saving", "Maturity", "Verdict", "Gaps"]
        for j, h in enumerate(headers):
            tbl.rows[0].cells[j].text = h
        for i, row in enumerate(sme_rows):
            tbl.rows[i + 1].cells[0].text = str(row.get("category") or "")
            tbl.rows[i + 1].cells[1].text = str(row.get("lever") or "")
            tbl.rows[i + 1].cells[2].text = str(row.get("saving_3yr") or "")
            tbl.rows[i + 1].cells[3].text = str(row.get("maturity") or "")
            tbl.rows[i + 1].cells[4].text = str(row.get("verdict") or "")
            tbl.rows[i + 1].cells[5].text = str(row.get("gaps") or "")
    else:
        doc.add_paragraph("No SME critique output.")

    doc.add_heading("Benchmark provenance", level=1)
    bench_rows = appendix.get("benchmark_rows") or []
    if bench_rows:
        tbl = doc.add_table(rows=1 + len(bench_rows), cols=6)
        tbl.style = "Table Grid"
        headers = ["Category", "Actual %", "P50 %", "Band", "Gap", "Vintage / n"]
        for j, h in enumerate(headers):
            tbl.rows[0].cells[j].text = h
        for i, row in enumerate(bench_rows):
            tbl.rows[i + 1].cells[0].text = str(row.get("category") or "")
            tbl.rows[i + 1].cells[1].text = str(row.get("actual_pct") or "")
            tbl.rows[i + 1].cells[2].text = str(row.get("p50_pct") or "")
            tbl.rows[i + 1].cells[3].text = str(row.get("band") or "")
            tbl.rows[i + 1].cells[4].text = str(row.get("gap") or "")
            tbl.rows[i + 1].cells[5].text = f"{row.get('vintage') or '—'} / n={row.get('sample_n') or '—'}"
    else:
        doc.add_paragraph("No benchmark comparisons.")

    doc.add_heading("Ingestion", level=1)
    ingestion = appendix.get("ingestion_summary")
    doc.add_paragraph(str(ingestion) if ingestion else "No ingestion report on file.")

    doc.save(path)
    return path
