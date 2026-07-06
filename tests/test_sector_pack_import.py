"""Tests for sector pack Excel import."""
from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from app.services.sector_pack_authoring import import_pack_from_workbook


def _build_template_workbook() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    overview = wb.create_sheet("1. Overview")
    overview.append(["pack_id", "sector", "version", "status"])
    overview.append([])
    overview.append([])
    overview.append([
        "import_test",
        "Import Test Sector",
        "0.1",
        "draft",
        "listed_company",
        "RBI",
        14.0,
        25.0,
        "March 31",
        "INR",
        "2026-01-01",
    ])

    levers = wb.create_sheet("2. Sector Levers")
    levers.append(["pack_id", "lever_id", "lever_name", "category_ids"])
    levers.append([])
    levers.append([])
    levers.append([
        "import_test",
        "imp_lever",
        "Import Lever",
        "IMPORT_CORE",
        1.0,
        3.0,
        5.0,
        "run_rate",
        10,
        "",
        "",
        "stub",
        "Test lever",
    ])

    tax = wb.create_sheet("6. Taxonomy Extensions")
    tax.append(["pack_id", "type", "category_id", "category_name"])
    tax.append([])
    tax.append([])
    tax.append(["import_test", "Additional", "IMPORT_CORE", "Import Core", "IMP", "operating"])

    peers = wb.create_sheet("4. Peer Benchmarks")
    peers.append(["pack_id", "name", "ticker"])
    peers.append([])
    peers.append([])
    peers.append(["import_test", "Import Peer Co", "IMP.NS", "NSE", "listed", 1000])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_workbook_dry_run():
    data = _build_template_workbook()
    result = import_pack_from_workbook(data, dry_run=True, pack_id="import_test")
    assert result["dry_run"] is True
    assert "import_test" in result["packs"]
    pack = result["packs"]["import_test"]
    assert "manifest" in pack["components"]
    assert pack["errors"] == []


def test_import_workbook_invalid_pack_id():
    data = _build_template_workbook()
    from app.services.sector_pack_authoring import SectorPackAuthoringError

    with pytest.raises(SectorPackAuthoringError):
        import_pack_from_workbook(data, dry_run=True, pack_id="missing_pack")
