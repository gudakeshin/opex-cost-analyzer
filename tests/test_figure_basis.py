"""Exports state what their figures assume and which caveats apply."""
from __future__ import annotations

from docx import Document

from app.services.board_deck import build_board_deck
from app.services.cfo_brief import build_cfo_brief, export_cfo_brief_docx
from app.services.evidence_appendix import build_evidence_appendix, export_evidence_appendix_docx
from app.services.figure_basis import basis_lines, basis_of_figures


def _analysis(**over):
    a = {
        "reporting_currency": "INR",
        "skill_outputs": {
            "savings-modeler": {
                "initiatives": [],
                "summary": {"total_expected_net_3yr": 6.0e7, "total_net_3yr_full_delivery": 8.0e7},
            },
            "peer-benchmarker": {
                "warnings": ["Annual revenue is missing or zero, so peer savings are not computed."],
                "benchmark_metadata": {"is_illustrative": True},
            },
            "value-bridge-calculator": {"confidence_bands": {"low": 1e7, "mid": 2e7, "high": 3e7}},
        },
        "ingestion_report": {"warnings": ["12 of 100 rows were dropped because the amount could not be read."]},
    }
    a.update(over)
    return a


def test_basis_states_expected_vs_full_delivery_and_caveats():
    lines = basis_lines(basis_of_figures(_analysis()))
    assert "after execution risk" in lines[0] and "6.00 Cr" in lines[0] and "8.00 Cr" in lines[0]
    assert any("revenue is missing" in l for l in lines)
    assert any("illustrative" in l for l in lines)
    assert any("rows were dropped" in l for l in lines)


def test_basis_is_empty_when_there_is_nothing_to_say():
    assert basis_lines(basis_of_figures({"skill_outputs": {}})) == []


def test_exports_include_the_basis(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.cfo_brief.OUTPUT_DIR", tmp_path)
    monkeypatch.setattr("app.services.evidence_appendix.OUTPUT_DIR", tmp_path)
    analysis = _analysis()

    brief = build_cfo_brief(analysis, company_name="Acme")
    assert brief["sections"]["basis_of_figures"]
    text = "\n".join(p.text for p in Document(export_cfo_brief_docx(brief, "b.docx")).paragraphs)
    assert "Basis of figures" in text and "after execution risk" in text

    deck = build_board_deck(analysis, company_name="Acme")
    assert "after execution risk" in deck["slides"][-1]["content"]["basis_of_figures"]

    appendix = build_evidence_appendix(analysis, company_name="Acme")
    atext = "\n".join(p.text for p in Document(export_evidence_appendix_docx(appendix, "e.docx")).paragraphs)
    assert "Basis of figures" in atext and "illustrative" in atext
