#!/usr/bin/env python3
"""Regenerate the data rows of the two hand-distributed pack workbooks from the
current sector_packs/ content, so the xlsx deliverables track the repo:

  - Sector_Packs_Documentation.xlsx (Overview, Sector Levers, KPI Packs,
    Peer Benchmarks, Benchmark Sources, Taxonomy Extensions)
  - Sector_Benchmarks_and_Levers.xlsx (Summary, Levers by Sector,
    Benchmarks by Sector)

Titles and header rows (rows 1-3) are preserved; data rows are rewritten.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
PACKS = ROOT / "sector_packs"

BASE_CATEGORY_NAMES = {
    "HR": "Human Resources (Employee Cost)",
    "IT": "Information Technology",
    "IT_SOFTWARE": "IT & Software",
    "FACILITIES": "Facilities & Real Estate",
    "MARKETING": "Marketing & Advertising",
    "TRAVEL": "Travel & Entertainment",
    "OUTSOURCED": "Outsourced / Contract Services",
    "PROF_SVCS": "Professional Services",
    "PROF_SERVICES": "Professional Services",
    "OFFICE": "Office & Administration",
    "TELECOM": "Telecom & Connectivity",
    "INSURANCE": "Insurance",
    "CONTINGENT": "Contingent Workforce",
    "LOGISTICS": "Logistics & Freight",
    "LOGISTICS_INDIA": "Logistics & Freight (India)",
    "FIN_SVCS": "Financial Services & Bank Charges",
    "RAW_MATERIALS": "Raw Materials",
    "CONVERSION_COST": "Conversion Cost",
    "PACKAGING": "Packaging Materials",
    "PLANT_MAINTENANCE": "Plant Maintenance & MRO",
    "POWER_ENERGY": "Power & Energy",
    "RND": "Research & Development",
    "BRSR_ESG": "BRSR / ESG Compliance",
    "FUEL_GENERATION": "Fuel & Generation",
    "GRID_MAINTENANCE": "Grid Maintenance",
    "T&D_LOSSES": "Transmission & Distribution Losses",
    "NETWORK_OPEX": "Network Operating Expenses",
    "SPECTRUM_LICENSING": "Spectrum & Licensing",
    "TOWER_INFRA": "Tower Infrastructure",
}


def load_json(p: Path):
    return json.loads(p.read_text(encoding="utf-8"))


def load_pack(pid: str) -> dict:
    d = PACKS / pid
    return {
        "manifest": yaml.safe_load((d / "pack_manifest.yaml").read_text(encoding="utf-8")),
        "taxonomy": load_json(d / "taxonomy_extension.json"),
        "peers": load_json(d / "peer_set.json"),
        "kpis": load_json(d / "kpi_pack.json"),
        "bench": load_json(d / "benchmarks_percentiles.json"),
        "levers": load_json(d / "sector_levers.json"),
        "sources": yaml.safe_load((d / "benchmark_sources.yaml").read_text(encoding="utf-8")),
    }


def pack_ids() -> list[str]:
    return sorted(p.name for p in PACKS.iterdir() if (p / "pack_manifest.yaml").exists())


def cat_name(pid: str, packs: dict, cat_id: str) -> str:
    for c in packs[pid]["taxonomy"].get("additional_categories", []):
        if c["category_id"] == cat_id:
            return c.get("category_name", cat_id)
    return BASE_CATEGORY_NAMES.get(cat_id, cat_id)


def clear_data_rows(ws, first_data_row: int = 4) -> None:
    if ws.max_row >= first_data_row:
        ws.delete_rows(first_data_row, ws.max_row - first_data_row + 1)


def refresh_documentation(packs: dict) -> None:
    path = ROOT / "Sector_Packs_Documentation.xlsx"
    wb = openpyxl.load_workbook(path)

    ws = wb["1. Overview"]
    clear_data_rows(ws)
    for pid in pack_ids():
        m = packs[pid]["manifest"]
        ws.append([
            pid, m.get("sector", ""), str(m.get("version", "")),
            str(m.get("status", "")).title(),
            ", ".join(map(str, m.get("applicable_entity_types", []))),
            ", ".join(map(str, m.get("regulatory_bodies", []))),
            m.get("default_wacc_pct", ""), m.get("default_tax_rate_pct", ""),
            m.get("fiscal_year_end", ""), m.get("reporting_currency", ""),
            str(m.get("created_at", "")),
        ])

    ws = wb["2. Sector Levers"]
    clear_data_rows(ws)
    for pid in pack_ids():
        for lv in packs[pid]["levers"].get("levers", []):
            ws.append([
                pid, lv.get("lever_id", ""), lv.get("lever_name", ""),
                ", ".join(lv.get("category_ids", [])),
                lv.get("p10_pct", ""), lv.get("p50_pct", ""), lv.get("p90_pct", ""),
                lv.get("savings_type", ""), lv.get("implementation_weeks", ""),
                lv.get("brsr_principle", ""),
                packs[pid]["levers"].get("source_type", ""),
                "; ".join((lv.get("industry_sources") or [])[:2]),
                lv.get("description", ""),
            ])

    ws = wb["3. KPI Packs"]
    clear_data_rows(ws)
    for pid in pack_ids():
        for k in packs[pid]["kpis"]:
            p50 = k.get("p50_benchmark_pct", "")
            ws.append([
                pid, k.get("kpi_id", ""), k.get("kpi_name", ""), k.get("unit", ""),
                k.get("formula", ""), str(k.get("target_direction", "")).title(),
                p50, k.get("unit", ""), k.get("source", ""),
            ])

    ws = wb["4. Peer Benchmarks"]
    clear_data_rows(ws)
    for pid in pack_ids():
        for peer in packs[pid]["peers"].get("peers", []):
            size = peer.get("revenue_cr") or peer.get("gwp_cr") or peer.get("group_revenue_cr") or peer.get("total_assets_cr") or ""
            ws.append([
                pid, peer.get("name", ""), peer.get("ticker", ""), peer.get("exchange", ""),
                peer.get("entity_type", ""), size, "", "", "", "",
            ])

    ws = wb["5. Benchmark Sources"]
    clear_data_rows(ws)
    for pid in pack_ids():
        src = packs[pid]["sources"] or {}
        for s in src.get("free_sources", []) or []:
            ws.append([
                pid, "Free", s.get("source_name") or s.get("name", ""),
                s.get("url", ""), ", ".join(map(str, s.get("data_available", []))) if isinstance(s.get("data_available"), list) else s.get("data_available", ""),
                s.get("refresh_frequency", src.get("refresh_cycle", "")),
            ])
        for s in src.get("licensed_sources_stubs", []) or []:
            ws.append([
                pid, "Licensed (stub)", s.get("source_name") or s.get("name", ""),
                s.get("url", ""), ", ".join(map(str, s.get("data_available", []))) if isinstance(s.get("data_available"), list) else s.get("data_available", ""),
                s.get("refresh_frequency", ""),
            ])

    ws = wb["6. Taxonomy Extensions"]
    clear_data_rows(ws)
    for pid in pack_ids():
        tax = packs[pid]["taxonomy"]
        for c in tax.get("additional_categories", []):
            ws.append([pid, "Additional", c.get("category_id", ""), c.get("category_name", ""),
                       c.get("gl_prefix", ""), c.get("cost_nature", "")])
        for c in tax.get("overrides", []) or []:
            ws.append([pid, "Override", c.get("category_id", ""), c.get("category_name", ""),
                       c.get("gl_prefix", ""), c.get("cost_nature", "")])

    wb.save(path)
    print(f"Refreshed {path.name}")


def spread_band(p25: float, p50: float, p90: float) -> str:
    try:
        ratio = (float(p90) - float(p25)) / max(float(p50), 0.01)
    except (TypeError, ValueError):
        return ""
    if ratio > 1.5:
        return "Wide"
    if ratio > 0.75:
        return "Moderate"
    return "Narrow"


def refresh_benchmarks_levers(packs: dict) -> None:
    path = ROOT / "Sector_Benchmarks_and_Levers.xlsx"
    wb = openpyxl.load_workbook(path)

    ws = wb["Summary"]
    clear_data_rows(ws)
    for pid in pack_ids():
        m = packs[pid]["manifest"]
        levers = packs[pid]["levers"].get("levers", [])
        bench = packs[pid]["bench"].get("benchmarks", {}).get(pid, {}).get("categories", {})
        p50s = [lv.get("p50_pct") for lv in levers if isinstance(lv.get("p50_pct"), (int, float))]
        top = max(levers, key=lambda lv: lv.get("p50_pct") or 0) if levers else {}
        ws.append([
            m.get("sector", pid), pid, str(m.get("status", "")).title(),
            len(levers), len(bench),
            round(sum(p50s) / len(p50s), 1) if p50s else "",
            top.get("lever_name", ""),
        ])

    ws = wb["Levers by Sector"]
    clear_data_rows(ws)
    for pid in pack_ids():
        m = packs[pid]["manifest"]
        for lv in packs[pid]["levers"].get("levers", []):
            ws.append([
                m.get("sector", pid), str(m.get("status", "")).title(),
                lv.get("lever_name", ""), lv.get("lever_id", ""), lv.get("description", ""),
                lv.get("p10_pct", ""), lv.get("p50_pct", ""), lv.get("p90_pct", ""),
                lv.get("implementation_weeks", ""),
                str(lv.get("savings_type", "")).replace("_", " ").title(),
                ", ".join(lv.get("category_ids", [])),
            ])

    ws = wb["Benchmarks by Sector"]
    clear_data_rows(ws)
    for pid in pack_ids():
        m = packs[pid]["manifest"]
        bench = packs[pid]["bench"].get("benchmarks", {}).get(pid, {}).get("categories", {})
        for cid, row in bench.items():
            ws.append([
                m.get("sector", pid), str(m.get("status", "")).title(), cid,
                cat_name(pid, packs, cid),
                row.get("P25", ""), row.get("P50", ""), row.get("P75", ""), row.get("P90", ""),
                spread_band(row.get("P25"), row.get("P50"), row.get("P90")),
            ])

    wb.save(path)
    print(f"Refreshed {path.name}")


def main() -> None:
    packs = {pid: load_pack(pid) for pid in pack_ids()}
    refresh_documentation(packs)
    refresh_benchmarks_levers(packs)


if __name__ == "__main__":
    main()
